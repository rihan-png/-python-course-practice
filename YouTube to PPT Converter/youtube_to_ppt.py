#!/usr/bin/env python3
"""
YouTube to PowerPoint Converter
================================
This script converts YouTube videos into PowerPoint presentations for teaching purposes.
It extracts the video transcript, key information, and creates slides that teachers can use.

Features:
- Extract video title and description
- Download and parse video transcript/subtitles
- Create PowerPoint slides with content
- Support for both automatic and manual captions
"""

import sys
import re
from typing import List, Dict, Optional
try:
    from youtube_transcript_api import YouTubeTranscriptApi
    from pytube import YouTube
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
except ImportError as e:
    print(f"Error: Missing required library - {e}")
    print("\nPlease install required packages:")
    print("pip install -r requirements.txt")
    sys.exit(1)


class YouTubeToPPTConverter:
    """Convert YouTube videos to PowerPoint presentations."""
    
    def __init__(self, video_url: str):
        """
        Initialize the converter with a YouTube URL.
        
        Args:
            video_url: YouTube video URL or video ID
        """
        self.video_url = video_url
        self.video_id = self._extract_video_id(video_url)
        self.video_info = None
        self.transcript = None
        self.prs = Presentation()
        
    def _extract_video_id(self, url: str) -> str:
        """
        Extract video ID from YouTube URL.
        
        Args:
            url: YouTube URL or video ID
            
        Returns:
            Video ID string
        """
        # If it's already a video ID (11 characters)
        if len(url) == 11 and not ('/' in url or '.' in url):
            return url
            
        # Extract from various YouTube URL formats
        patterns = [
            r'(?:youtube\.com\/watch\?v=|youtu\.be\/)([^&\n?]*)',
            r'youtube\.com\/embed\/([^&\n?]*)',
            r'youtube\.com\/v\/([^&\n?]*)'
        ]
        
        for pattern in patterns:
            match = re.search(pattern, url)
            if match:
                return match.group(1)
                
        raise ValueError(f"Could not extract video ID from URL: {url}")
    
    def fetch_video_info(self) -> Dict:
        """
        Fetch video information using pytube.
        
        Returns:
            Dictionary containing video metadata
        """
        try:
            yt = YouTube(self.video_url)
            self.video_info = {
                'title': yt.title,
                'author': yt.author,
                'length': yt.length,
                'views': yt.views,
                'description': yt.description or "No description available"
            }
            return self.video_info
        except Exception as e:
            print(f"Error fetching video info: {e}")
            # Fallback to basic info
            self.video_info = {
                'title': 'YouTube Video',
                'author': 'Unknown',
                'description': 'No description available'
            }
            return self.video_info
    
    def fetch_transcript(self) -> List[Dict]:
        """
        Fetch video transcript using YouTube Transcript API.
        
        Returns:
            List of transcript segments
        """
        try:
            transcript_list = YouTubeTranscriptApi.get_transcript(self.video_id)
            self.transcript = transcript_list
            return transcript_list
        except Exception as e:
            print(f"Error fetching transcript: {e}")
            print("Note: Some videos may not have transcripts available.")
            return []
    
    def _group_transcript_into_chunks(self, max_duration: int = 60) -> List[Dict]:
        """
        Group transcript into logical chunks based on time.
        
        Args:
            max_duration: Maximum duration for each chunk in seconds
            
        Returns:
            List of grouped transcript chunks
        """
        if not self.transcript:
            return []
        
        chunks = []
        current_chunk = {
            'start': 0,
            'text': '',
            'duration': 0
        }
        
        for entry in self.transcript:
            # If adding this entry would exceed max duration, start new chunk
            if current_chunk['duration'] + entry['duration'] > max_duration and current_chunk['text']:
                chunks.append(current_chunk)
                current_chunk = {
                    'start': entry['start'],
                    'text': entry['text'],
                    'duration': entry['duration']
                }
            else:
                if not current_chunk['text']:
                    current_chunk['start'] = entry['start']
                current_chunk['text'] += ' ' + entry['text']
                current_chunk['duration'] += entry['duration']
        
        # Add the last chunk
        if current_chunk['text']:
            chunks.append(current_chunk)
        
        return chunks
    
    def create_title_slide(self):
        """Create the title slide with video information."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[0])
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = self.video_info.get('title', 'YouTube Video')
        subtitle.text = f"By: {self.video_info.get('author', 'Unknown')}\nSource: YouTube"
    
    def create_content_slide(self, title_text: str, content_text: str, slide_number: int = None):
        """
        Create a content slide with title and body text.
        
        Args:
            title_text: Title for the slide
            content_text: Main content text
            slide_number: Optional slide number for title
        """
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[1])
        title = slide.shapes.title
        body = slide.placeholders[1]
        
        if slide_number:
            title.text = f"{title_text} - Part {slide_number}"
        else:
            title.text = title_text
        
        # Clean up the text
        content_text = content_text.strip()
        
        # Add content to text frame
        text_frame = body.text_frame
        text_frame.text = content_text
        text_frame.word_wrap = True
        
        # Format text
        for paragraph in text_frame.paragraphs:
            paragraph.font.size = Pt(18)
            paragraph.level = 0
    
    def create_description_slide(self):
        """Create a slide with video description."""
        description = self.video_info.get('description', 'No description available')
        
        # Limit description length for slide
        if len(description) > 500:
            description = description[:500] + "..."
        
        self.create_content_slide("About This Video", description)
    
    def generate_presentation(self, output_file: str = "youtube_presentation.pptx") -> str:
        """
        Generate the complete PowerPoint presentation.
        
        Args:
            output_file: Output filename for the presentation
            
        Returns:
            Path to the generated presentation
        """
        print(f"Generating presentation for video: {self.video_id}")
        
        # Fetch video information
        print("Fetching video information...")
        self.fetch_video_info()
        
        # Create title slide
        print("Creating title slide...")
        self.create_title_slide()
        
        # Create description slide
        print("Creating description slide...")
        self.create_description_slide()
        
        # Fetch and process transcript
        print("Fetching transcript...")
        transcript = self.fetch_transcript()
        
        if transcript:
            print(f"Found {len(transcript)} transcript entries")
            
            # Group transcript into logical chunks
            chunks = self._group_transcript_into_chunks(max_duration=90)
            print(f"Created {len(chunks)} content chunks")
            
            # Create content slides from transcript chunks
            for i, chunk in enumerate(chunks, 1):
                # Format time for title
                minutes = int(chunk['start'] // 60)
                seconds = int(chunk['start'] % 60)
                time_str = f"{minutes}:{seconds:02d}"
                
                title = f"Content at {time_str}"
                self.create_content_slide(title, chunk['text'], i)
        else:
            # Create a placeholder slide if no transcript
            self.create_content_slide(
                "No Transcript Available",
                "This video does not have an available transcript.\n\n"
                "You can manually add content to these slides based on the video."
            )
        
        # Save presentation
        print(f"Saving presentation to {output_file}...")
        self.prs.save(output_file)
        print(f"✓ Presentation created successfully: {output_file}")
        
        return output_file


def main():
    """Main function to run the converter from command line."""
    print("=" * 60)
    print("YouTube to PowerPoint Converter")
    print("=" * 60)
    print()
    
    if len(sys.argv) < 2:
        print("Usage: python youtube_to_ppt.py <youtube_url> [output_file]")
        print()
        print("Examples:")
        print("  python youtube_to_ppt.py https://www.youtube.com/watch?v=VIDEO_ID")
        print("  python youtube_to_ppt.py VIDEO_ID my_presentation.pptx")
        print()
        sys.exit(1)
    
    video_url = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else "youtube_presentation.pptx"
    
    try:
        converter = YouTubeToPPTConverter(video_url)
        converter.generate_presentation(output_file)
        print()
        print("✓ Done! You can now use this presentation for teaching.")
        
    except Exception as e:
        print(f"\n✗ Error: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
